"""Write the compositions workbook, beside the customer's documents, once they have approved it.

Usage:
    python3 write_workbook.py <session_dir> <documents_folder> [--date YYYY-MM-DD]

Reads <session_dir>/COMPOSITIONS.json and writes
<documents_folder>/RAW_MATERIAL_COMPOSITIONS <date>.xlsx.

Dated rather than overwritten, because the customer's folder is an input and a second run must not
quietly replace the file somebody has already started working from.

Three sheets:

  Composition            one row per ingredient, with the line it was read off
  Reference Documents    every file the run considered per material, and why each one lost
  MATERIALS_WITH_ISSUES  every material carrying no composition, and what stopped it

The middle sheet is the one that answers "why this file and not that one". A workbook that records only
the file it used cannot answer it, and that is the question a person asks first.
"""

import argparse
import json
import sys
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))

import owed

HEAD = PatternFill("solid", fgColor="1F3864")
TITLE = Font(bold=True, color="FFFFFF")


def sheet(wb, title, headers, widths, rows, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    for i, header in enumerate(headers, 1):
        cell = ws.cell(1, i, header)
        cell.font = TITLE
        cell.fill = HEAD
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for n, row in enumerate(rows, 2):
        for i, value in enumerate(row, 1):
            ws.cell(n, i, value).alignment = Alignment(vertical="top", wrap_text=False)
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{1 + len(rows)}"
    return ws


def best_date(dates):
    """The latest date a candidate prints: the words on the page, and the kind of date they label.

    Only ever used for a file that lost. The winner's date is the one its sub agent read and judged,
    which is a different and better thing than the largest number on the page.
    """
    resolved = [entry for entry in dates or [] if entry.get("iso")]
    if not resolved:
        return "", ""
    latest = max(resolved, key=lambda entry: entry["iso"])
    return latest.get("written", ""), latest.get("kind", "")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("session_dir", type=Path)
    parser.add_argument("documents_folder", type=Path)
    parser.add_argument("--date", default=date.today().isoformat())
    args = parser.parse_args()

    state = json.loads((args.session_dir / "COMPOSITIONS.json").read_text(encoding="utf-8"))
    materials = state["materials"]

    composition = [
        (m["identifier"], m["itemName"], row["ingredient"], row["cas"], row["ec"],
         row["min"], row["max"], row["quote"], m["sourceFile"],
         "; ".join(m.get("sourceFolders") or []))
        for m in sorted(materials, key=lambda m: m["identifier"]) for row in m["rows"]]

    references = []
    for material in sorted(materials, key=lambda m: m["identifier"]):
        printed = material.get("documentDate") or {}
        for entry in material["considered"]:
            references.append((
                material["identifier"], entry["fileName"], entry["folder"],
                "PRIMARY" if entry["chosen"] else "considered",
                *((printed.get("written", ""), printed.get("kind", "")) if entry["chosen"]
                  else best_date(entry.get("dates"))),
                material["why"] if entry["chosen"] else entry["why"]))

    issues = [(issue["identifier"], issue.get("sourceFile", ""),
               "; ".join(issue.get("candidates") or []), issue["why"])
              for issue in sorted(state["issues"], key=lambda i: i["identifier"])]

    # A material can carry a composition and still owe somebody a look. These sit on the issues tab
    # beside the materials that carry nothing, because the action is the same: read the page. The
    # ingredient goes in the last column with the reason - the third is "files considered", and a row
    # is not a file.
    incomplete = 0
    for material in sorted(materials, key=lambda m: m["identifier"]):
        for name, why in owed.owed_on(material):
            issues.append((material["identifier"], material["sourceFile"], "",
                           f"{name}: {why}" if name else why))
        incomplete += bool(owed.owed_on(material))

    wb = openpyxl.Workbook()
    sheet(wb, "Composition",
          ["code", "raw material", "ingredient", "CAS", "EC", "min %", "max %",
           "quoted line", "source file", "source folders"],
          [12, 30, 40, 16, 14, 9, 9, 60, 46, 40], composition, first=True)
    sheet(wb, "Reference Documents",
          ["code", "file", "source folder", "role", "date printed", "kind of date", "why"],
          [12, 50, 34, 13, 16, 13, 70], references)
    sheet(wb, "MATERIALS_WITH_ISSUES",
          ["code", "file it tried", "files considered", "what stopped it"],
          [12, 44, 56, 80], issues)

    out = args.documents_folder / f"RAW_MATERIAL_COMPOSITIONS {args.date}.xlsx"
    # Two approvals in one day are two workbooks. The date alone is not a unique name, and somebody may
    # already be working from the one written an hour ago.
    run = 2
    while out.exists():
        out = args.documents_folder / f"RAW_MATERIAL_COMPOSITIONS {args.date} ({run}).xlsx"
        run += 1
    wb.save(out)

    print(f"{len(composition)} ingredient row(s) across {len(materials)} material(s)")
    print(f"  {len(references)} reference row(s), {len(state['issues'])} material(s) carrying nothing")
    if incomplete:
        print(f"  {incomplete} material(s) carry a composition and still owe somebody a look")
    if state.get("contradictions"):
        print(f"  {len(state['contradictions'])} document(s) read two different ways - see "
              f"COMPOSITIONS.json")
    print(f"\n-> {out}")


if __name__ == "__main__":
    main()
