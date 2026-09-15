"""The composition round trip: what reaches a sub agent, and what its answer has to survive.

The sub agent itself is not tested here — the contract either side of it is. Going in: a shortlist of
files that look like compositions, with the dates and codes each one prints. Coming back: rows that are
only carried if every value in them is in the document the answer names.
"""

import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
SKILL = REPO / "plugins/customer-service-skills/skills/extract-composition-from-documents"
LIB = SKILL / "lib"

COMPOSITION = """SAFETY DATA SHEET
Dissolvine GL-47-S
Issue Date: 14/03/2023   Print Date: 02/09/2024

SECTION 3 Composition / information on ingredients
INCI name                        CAS         EC          Concentration   Function
Tetrasodium glutamate diacetate  51981-21-6  257-573-7   47.5-48.0       Chelating agent
Aqua/water                       7732-18-5   231-791-2   balance         Solvent

SECTION 11 Toxicological information
Tetrasodium glutamate diacetate is not classified.
"""

HAZARD_BANDS = """SAFETY DATA SHEET
Issue Date: 01/06/2024
SECTION 3: Composition and information on ingredients
Name                 Product identifier      %
Hydantoin            CAS-No.: 6440-58-0      >= 30 - < 70
Methanol             CAS-No.: 67-56-1        >= 0.3 - < 0.5
"""

BROCHURE = """Marketing brochure
Our chelates keep your formulation stable. Ask about 100% renewable sourcing.
"""


class Session:
    """A session directory holding what the two passes before this one would have left."""

    def __init__(self, root):
        self.root = Path(root)
        (self.root / "extracted").mkdir(parents=True, exist_ok=True)
        self.documents, self.extracted = [], []
        (self.root / "BRANCHES.json").write_text(json.dumps({"branches": [
            {"name": "raw materials", "pathPrefix": "", "table": "raw_materials",
             "identifier": {"type": "folderLevelPattern", "level": 1,
                            "pattern": "^(R[0-9]+)"}}]}), encoding="utf-8")
        (self.root / "WORKFLOW.json").write_text(json.dumps({
            "documentTemplates": [{"id": "dt_sds", "name": "Safety Data Sheet (SDS)",
                                   "for_tables": ["raw_materials"]}],
            "itemTemplates": [{"name": "Raw Material", "table": "raw_materials",
                               "identifierColumn": "code", "documentSections": []}]}), encoding="utf-8")
        (self.root / "ITEMS.csv").write_text(
            "table,id,identifier,name,template,archived\n"
            "raw_materials,1,R289,Dissolvine GL-47-S,Raw Material,false\n", encoding="utf-8")

    def add(self, folder, name, body, sha=None):
        sha = sha or f"{len(self.documents):064d}"
        text_file = self.root / "extracted" / f"{sha}.md"
        text_file.write_text(body, encoding="utf-8")
        path = f"C:/folder/{folder}/{name}"
        self.documents.append({"path": path, "relativePath": f"{folder}/{name}",
                               "folder": folder, "name": name, "sha": sha})
        self.extracted.append({"path": path, "kind": "text",
                               "textFile": str(text_file).replace("\\", "/"),
                               "ocrTextFile": None, "images": []})
        return sha

    def close(self):
        (self.root / "DOCUMENTS.json").write_text(json.dumps(self.documents), encoding="utf-8")
        (self.root / "EXTRACTED.json").write_text(json.dumps({"documents": self.extracted}),
                                                  encoding="utf-8")

    def run(self, script, *arguments):
        result = subprocess.run([sys.executable, str(LIB / script), str(self.root), *arguments],
                                capture_output=True, text=True)
        if result.returncode:
            raise AssertionError(f"{script} failed:\n{result.stdout}\n{result.stderr}")
        return result.stdout

    def read(self, name):
        return json.loads((self.root / name).read_text(encoding="utf-8"))

    def answer(self, identifier, body):
        (self.root / "read").mkdir(parents=True, exist_ok=True)
        (self.root / "read" / f"{identifier}.json").write_text(json.dumps(body), encoding="utf-8")


def good_answer(sha):
    return {
        "identifier": "R289", "sourceFile": "TDS.pdf", "sourceSha": sha,
        "documentDate": {"written": "14/03/2023", "label": "Issue Date", "kind": "issue"},
        "why": "the only file with a composition table",
        "rows": [
            {"ingredient": "Tetrasodium Glutamate Diacetate", "cas": "51981-21-6", "ec": "257-573-7",
             "min": 47.5, "max": 48.0, "statedAs": "47.5-48.0",
             "quote": "Tetrasodium glutamate diacetate  51981-21-6  257-573-7   47.5-48.0"},
            {"ingredient": "Aqua", "cas": "7732-18-5", "ec": "231-791-2",
             "min": None, "max": None, "statedAs": "balance",
             "quote": "Aqua/water                       7732-18-5   231-791-2   balance"},
        ],
    }


class ShortlistTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.session = Session(self.tmp.name)
        self.sha = self.session.add("R289 - Dissolvine", "TDS.pdf", COMPOSITION)
        self.session.add("R289 - Dissolvine", "brochure.pdf", BROCHURE)
        self.session.close()

    def tearDown(self):
        self.tmp.cleanup()

    def test_a_composition_is_shortlisted_and_a_brochure_is_not(self):
        self.session.run("score_composition.py")
        material = self.session.read("CANDIDATES.json")["materials"][0]
        self.assertEqual(material["identifier"], "R289")
        self.assertEqual([c["fileName"] for c in material["candidates"]], ["TDS.pdf"])

    def test_the_task_carries_the_dates_and_codes_the_agent_must_not_invent(self):
        self.session.run("score_composition.py")
        self.session.run("plan_readings.py")
        task = json.loads((self.session.root / "readings" / "R289.json").read_text(encoding="utf-8"))
        entry = task["files"][0]
        self.assertIn("51981-21-6", entry["casNumbers"])
        self.assertIn("231-791-2", entry["ecNumbers"])
        printed = {(d["written"], d["kind"]) for d in entry["datesPrinted"]}
        self.assertIn(("14/03/2023", "issue"), printed)
        self.assertIn(("02/09/2024", "print"), printed)


class VerificationTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.session = Session(self.tmp.name)
        self.sha = self.session.add("R289 - Dissolvine", "TDS.pdf", COMPOSITION)
        self.other = self.session.add("R289 - Dissolvine", "SDS.pdf", HAZARD_BANDS)
        self.session.close()
        self.session.run("score_composition.py")
        self.session.run("plan_readings.py")

    def tearDown(self):
        self.tmp.cleanup()

    def collect(self):
        self.session.run("collect_compositions.py")
        return self.session.read("COMPOSITIONS.json")

    def test_a_transcribed_composition_is_carried(self):
        self.session.answer("R289", good_answer(self.sha))
        state = self.collect()
        self.assertEqual(state["issues"], [])
        rows = state["materials"][0]["rows"]
        self.assertEqual([r["ingredient"] for r in rows],
                         ["Tetrasodium Glutamate Diacetate", "Aqua"])
        self.assertEqual(rows[1]["statedAs"], "balance")
        self.assertIsNone(rows[1]["min"], "a balance line states no figure and must not be given one")

    def test_an_invented_line_blocks_the_whole_material(self):
        answer = good_answer(self.sha)
        answer["rows"].append({"ingredient": "Sodium Hydroxide", "cas": "1310-73-2", "ec": "",
                               "min": 0.5, "max": 0.9,
                               "quote": "Sodium hydroxide  1310-73-2  0.5-0.9  By-product"})
        self.session.answer("R289", answer)
        state = self.collect()
        self.assertEqual(state["materials"], [])
        self.assertIn("not in the document", state["issues"][0]["why"])

    def test_a_cas_number_the_document_does_not_print_is_refused(self):
        answer = good_answer(self.sha)
        answer["rows"][1]["cas"] = "56-81-5"
        self.session.answer("R289", answer)
        state = self.collect()
        self.assertEqual(state["materials"], [])
        self.assertIn("is not printed in the document", state["issues"][0]["why"])

    def test_a_mistyped_cas_number_fails_its_check_digit(self):
        answer = good_answer(self.sha)
        answer["rows"][0]["cas"] = "51981-21-7"
        self.session.answer("R289", answer)
        state = self.collect()
        self.assertEqual(state["materials"], [])
        self.assertIn("check digit", state["issues"][0]["why"])

    def test_a_date_the_document_does_not_print_is_refused(self):
        answer = good_answer(self.sha)
        answer["documentDate"] = {"written": "01/01/2025", "label": "Issue Date", "kind": "issue"}
        self.session.answer("R289", answer)
        state = self.collect()
        self.assertEqual(state["materials"], [])
        self.assertIn("is not printed in the document", state["issues"][0]["why"])

    def test_a_file_that_was_never_shortlisted_is_refused(self):
        answer = good_answer("f" * 64)
        self.session.answer("R289", answer)
        state = self.collect()
        self.assertEqual(state["materials"], [])
        self.assertIn("not on its shortlist", state["issues"][0]["why"])

    def test_a_material_nobody_answered_for_says_so(self):
        state = self.collect()
        self.assertEqual(state["unanswered"], ["R289"])
        self.assertIn("no sub agent answered", state["issues"][0]["why"])

    def test_an_agent_that_finds_no_composition_reports_it_rather_than_guessing(self):
        self.session.answer("R289", {"identifier": "R289", "sourceSha": None,
                                     "why": "both files give hazard bands, neither states a composition"})
        state = self.collect()
        self.assertEqual(state["materials"], [])
        self.assertIn("hazard bands", state["issues"][0]["why"])

    def test_a_recut_narrows_the_task_to_the_file_the_user_named(self):
        self.session.answer("R289", good_answer(self.sha))
        self.collect()
        recut = self.session.root / "recut_R289.json"
        recut.write_text(json.dumps({"identifier": "R289", "sourceFile": "SDS.pdf",
                                     "why": "the newer sheet"}), encoding="utf-8")
        self.session.run("apply_recut.py", str(recut))
        task = json.loads((self.session.root / "readings" / "R289.json").read_text(encoding="utf-8"))
        self.assertEqual([f["fileName"] for f in task["files"]], ["SDS.pdf"])
        self.assertFalse((self.session.root / "read" / "R289.json").is_file(),
                         "the old answer has to go, or the collector carries it unchanged")


class WorkbookTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.session = Session(self.tmp.name)
        self.sha = self.session.add("R289 - Dissolvine", "TDS.pdf", COMPOSITION)
        self.session.close()
        self.session.run("score_composition.py")
        self.session.run("plan_readings.py")
        self.session.answer("R289", good_answer(self.sha))
        self.session.run("collect_compositions.py")

    def tearDown(self):
        self.tmp.cleanup()

    def test_the_workbook_is_dated_and_carries_three_sheets(self):
        folder = Path(self.tmp.name) / "documents"
        folder.mkdir()
        self.session.run("write_workbook.py", str(folder), "--date", "2026-09-15")
        out = folder / "RAW_MATERIAL_COMPOSITIONS 2026-09-15.xlsx"
        self.assertTrue(out.is_file())
        import openpyxl
        wb = openpyxl.load_workbook(out)
        self.assertEqual(wb.sheetnames,
                         ["Composition", "Reference Documents", "MATERIALS_WITH_ISSUES"])
        self.assertEqual(wb["Composition"].max_row - 1, 2)


if __name__ == "__main__":
    unittest.main()
