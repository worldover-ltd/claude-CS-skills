"""What a reader does with a row it can see but cannot carry, and what a multi-code cell no longer costs.

The run blocks a whole material when any value fails, which is right when a value was invented and
punishing when a document simply prints three CAS numbers against one name. Faced with that, a reader
learns the cheapest safe move is to say nothing: on one real material it returned four of eighteen rows
and the omission was invisible - the four looked like a composition and summed near enough to 100.

So a row it cannot carry goes in `withheld` instead. Withheld rows never reach the workbook's
composition - nothing about them has been checked - and they never block the material either. They are
the run saying out loud that it saw more than it could take.
"""

import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
SKILL = REPO / "plugins/customer-service-skills/skills/extract-composition-from-documents"
LIB = SKILL / "lib"
sys.path.insert(0, str(LIB))

sys.path.insert(0, str(REPO / "tests"))
from test_composition import Session  # noqa: E402


ITEMS = """table,id,identifier,name,template,archived
raw_materials,1,R131,Tender Green Fragrance,Raw Material,false
"""


def session_for(root):
    """A session whose app holds R131, which the shared fixture's items file does not."""
    session = Session(root)
    (session.root / "ITEMS.csv").write_text(ITEMS, encoding="utf-8")
    return session


FRAGRANCE = """Tender Green Fragrance
Composition breakdown
Issue Date: 24/06/2024

Chemical Name                CAS RN                              TOTAL
Limonene                     5989-27-5; 138-86-3; 7705-14-8      38.5
Terpineol                    8000-41-7; 98-55-5; 138-87-4        7.0
Citronellol                  106-22-9; 1117-61-9                 0.06
"""


class AMultiCodeCellNoLongerBlocks(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.session = session_for(self.tmp.name)
        self.sha = self.session.add("R131 - Tender Green", "SB312.pdf", FRAGRANCE)
        self.session.close()
        self.session.run("score_composition.py")
        self.session.run("plan_readings.py")

    def answer(self, rows, withheld=None):
        body = {
            "identifier": "R131", "sourceSha": self.sha, "sourceFile": "SB312.pdf",
            "documentDate": {"written": "24/06/2024", "label": "Issue Date", "kind": "issue"},
            "why": "the composition breakdown", "rows": rows,
        }
        if withheld is not None:
            body["withheld"] = withheld
        self.session.answer("R131", body)
        self.session.run("collect_compositions.py")
        return self.session.read("COMPOSITIONS.json")

    def test_a_row_whose_cas_cell_holds_three_numbers_is_carried(self):
        state = self.answer([{
            "ingredient": "Limonene", "cas": "5989-27-5; 138-86-3; 7705-14-8", "ec": "",
            "min": 38.5, "max": 38.5,
            "quote": "Limonene                     5989-27-5; 138-86-3; 7705-14-8      38.5"}])
        self.assertEqual(len(state["materials"]), 1, state["issues"])
        self.assertEqual(state["materials"][0]["rows"][0]["cas"], "5989-27-5; 138-86-3; 7705-14-8")

    def test_a_mistyped_digit_in_such_a_cell_still_blocks(self):
        state = self.answer([{
            "ingredient": "Limonene", "cas": "5989-27-5; 138-86-4; 7705-14-8", "ec": "",
            "min": 38.5, "max": 38.5,
            "quote": "Limonene                     5989-27-5; 138-86-3; 7705-14-8      38.5"}])
        self.assertEqual(state["materials"], [])
        self.assertIn("check digit", state["issues"][0]["why"])

    def test_a_number_not_in_the_document_still_blocks(self):
        state = self.answer([{
            "ingredient": "Limonene", "cas": "5989-27-5; 110-82-7", "ec": "",
            "min": 38.5, "max": 38.5,
            "quote": "Limonene                     5989-27-5; 138-86-3; 7705-14-8      38.5"}])
        self.assertEqual(state["materials"], [])
        self.assertIn("not printed", state["issues"][0]["why"])


class WhatTheReaderCouldNotCarry(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.session = session_for(self.tmp.name)
        self.sha = self.session.add("R131 - Tender Green", "SB312.pdf", FRAGRANCE)
        self.session.close()
        self.session.run("score_composition.py")
        self.session.run("plan_readings.py")

    def answer(self, withheld):
        body = {
            "identifier": "R131", "sourceSha": self.sha, "sourceFile": "SB312.pdf",
            "documentDate": {"written": "24/06/2024", "label": "Issue Date", "kind": "issue"},
            "why": "the composition breakdown", "withheld": withheld,
            "rows": [{"ingredient": "Limonene", "cas": "5989-27-5", "ec": "",
                      "min": 38.5, "max": 38.5,
                      "quote": "Limonene                     5989-27-5; 138-86-3; 7705-14-8      38.5"}],
        }
        self.session.answer("R131", body)
        self.session.run("collect_compositions.py")
        return self.session.read("COMPOSITIONS.json")

    def test_a_withheld_row_does_not_block_the_material(self):
        state = self.answer([{"ingredient": "Terpineol", "asPrinted": "8000-41-7; 98-55-5; 138-87-4  7.0",
                              "why": "three CAS numbers and I cannot tell which is this row's"}])
        self.assertEqual(len(state["materials"]), 1, state["issues"])
        self.assertEqual(len(state["materials"][0]["rows"]), 1)

    def test_it_is_carried_beside_the_composition_with_its_reason(self):
        state = self.answer([{"ingredient": "Terpineol", "asPrinted": "8000-41-7; 98-55-5; 138-87-4  7.0",
                              "why": "three CAS numbers and I cannot tell which is this row's"}])
        withheld = state["materials"][0]["withheld"]
        self.assertEqual(withheld[0]["ingredient"], "Terpineol")
        self.assertIn("three CAS", withheld[0]["why"])

    def test_it_never_reaches_the_composition_rows(self):
        state = self.answer([{"ingredient": "Terpineol", "asPrinted": "8000-41-7; 98-55-5; 138-87-4  7.0",
                              "why": "unsure"}])
        self.assertNotIn("Terpineol", [r["ingredient"] for r in state["materials"][0]["rows"]])

    def test_a_line_it_names_is_still_held_against_the_document(self):
        """Not to refuse it - to say so. An invented withheld row should be as visible as an invented one."""
        state = self.answer([{"ingredient": "Invented", "asPrinted": "Invented  99-99-9  4.0",
                              "why": "unsure"}])
        self.assertEqual(len(state["materials"]), 1)
        self.assertFalse(state["materials"][0]["withheld"][0]["inDocument"])

    def test_a_withheld_row_that_is_on_the_page_says_so(self):
        state = self.answer([{"ingredient": "Terpineol",
                              "asPrinted": "Terpineol                    8000-41-7; 98-55-5; 138-87-4        7.0",
                              "why": "unsure"}])
        self.assertTrue(state["materials"][0]["withheld"][0]["inDocument"])

    def test_a_withheld_list_of_the_wrong_shape_costs_one_material_not_the_run(self):
        """Every other malformed answer is one material's issue; this one used to be everybody's."""
        state = self.answer({"Terpineol": "unsure"})
        self.assertEqual(len(state["materials"]), 1, state["issues"])
        self.assertEqual(state["materials"][0]["withheld"], [])

    def test_an_unprinted_ec_number_beside_a_good_cas_still_blocks(self):
        body = {
            "identifier": "R131", "sourceSha": self.sha, "sourceFile": "SB312.pdf",
            "documentDate": {"written": "24/06/2024", "label": "Issue Date", "kind": "issue"},
            "why": "the composition breakdown",
            "rows": [{"ingredient": "Limonene", "cas": "5989-27-5 / EC 231-791-2", "ec": "",
                      "min": 38.5, "max": 38.5,
                      "quote": "Limonene                     5989-27-5; 138-86-3; 7705-14-8      38.5"}],
        }
        self.session.answer("R131", body)
        self.session.run("collect_compositions.py")
        state = self.session.read("COMPOSITIONS.json")
        self.assertEqual(state["materials"], [])
        self.assertIn("231-791-2", state["issues"][0]["why"])

    def test_no_withheld_list_at_all_is_fine(self):
        body = {
            "identifier": "R131", "sourceSha": self.sha, "sourceFile": "SB312.pdf",
            "documentDate": {"written": "24/06/2024", "label": "Issue Date", "kind": "issue"},
            "why": "the composition breakdown",
            "rows": [{"ingredient": "Limonene", "cas": "5989-27-5", "ec": "", "min": 38.5, "max": 38.5,
                      "quote": "Limonene                     5989-27-5; 138-86-3; 7705-14-8      38.5"}],
        }
        self.session.answer("R131", body)
        self.session.run("collect_compositions.py")
        state = self.session.read("COMPOSITIONS.json")
        self.assertEqual(state["materials"][0]["withheld"], [])


class ItReachesTheWorkbook(unittest.TestCase):
    """The composition tab is the thing people act on, so what it leaves out has to land somewhere."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.session = session_for(self.tmp.name)
        self.sha = self.session.add("R131 - Tender Green", "SB312.pdf", FRAGRANCE)
        self.session.close()
        self.session.run("score_composition.py")
        self.session.run("plan_readings.py")
        self.session.answer("R131", {
            "identifier": "R131", "sourceSha": self.sha, "sourceFile": "SB312.pdf",
            "documentDate": {"written": "24/06/2024", "label": "Issue Date", "kind": "issue"},
            "why": "the composition breakdown",
            "withheld": [{"ingredient": "Terpineol", "asPrinted": "8000-41-7; 98-55-5; 138-87-4  7.0",
                          "why": "the percentage column did not survive the conversion"}],
            "rows": [{"ingredient": "Limonene", "cas": "5989-27-5; 138-86-3; 7705-14-8", "ec": "",
                      "min": 38.5, "max": 38.5,
                      "quote": "Limonene                     5989-27-5; 138-86-3; 7705-14-8      38.5"}]})
        self.session.run("collect_compositions.py")
        self.folder = Path(self.tmp.name) / "documents"
        self.folder.mkdir()
        self.session.run("write_workbook.py", str(self.folder), "--date", "2026-09-16")

    def book(self):
        import openpyxl

        return openpyxl.load_workbook(self.folder / "RAW_MATERIAL_COMPOSITIONS 2026-09-16.xlsx")

    def test_the_withheld_row_is_on_the_issues_tab(self):
        sheet = self.book()["MATERIALS_WITH_ISSUES"]
        rows = [[c.value for c in row] for row in sheet.iter_rows(min_row=2)]
        found = [row for row in rows if "Terpineol" in str(row[3] or "")]
        self.assertTrue(found, rows)
        self.assertIn("could not carry", found[0][3])

    def test_it_lands_under_the_heading_that_describes_it(self):
        """Column three is "files considered", and a row of a table is not a file."""
        sheet = self.book()["MATERIALS_WITH_ISSUES"]
        headers = [c.value for c in next(sheet.iter_rows(max_row=1))]
        self.assertEqual(headers[3], "what stopped it")
        rows = [[c.value for c in row] for row in sheet.iter_rows(min_row=2)]
        found = [row for row in rows if "Terpineol" in str(row[3] or "")]
        self.assertFalse(found[0][2], "the ingredient must not be filed as a file considered")

    def test_it_is_not_on_the_composition_tab(self):
        sheet = self.book()["Composition"]
        names = [sheet.cell(n, 3).value for n in range(2, sheet.max_row + 1)]
        self.assertIn("Limonene", names)
        self.assertNotIn("Terpineol", names)

    def test_the_multi_code_cell_reaches_the_sheet_as_the_document_printed_it(self):
        sheet = self.book()["Composition"]
        self.assertEqual(sheet.cell(2, 4).value, "5989-27-5; 138-86-3; 7705-14-8")


class WhatAMaterialStillOwes(unittest.TestCase):
    """One list, read by both the workbook and the review card, so the two cannot disagree."""

    def test_a_withheld_row_is_owed_with_its_reason(self):
        import owed

        said = owed.owed_on({"rows": [], "withheld": [
            {"ingredient": "Terpineol", "why": "the column is garbled", "inDocument": True}]})
        self.assertEqual(len(said), 1)
        self.assertEqual(said[0][0], "Terpineol")
        self.assertIn("garbled", said[0][1])

    def test_a_row_whose_line_is_not_on_the_page_says_so(self):
        import owed

        said = owed.owed_on({"rows": [], "withheld": [
            {"ingredient": "Invented", "why": "unsure", "inDocument": False}]})
        self.assertIn("not in the document", said[0][1])

    def test_a_reader_that_gave_no_reason_still_reports_the_row(self):
        import owed

        said = owed.owed_on({"rows": [], "withheld": [{"ingredient": "Terpineol"}]})
        self.assertIn("Terpineol", said[0][0])

    def test_a_complete_answer_owes_nothing(self):
        import owed

        self.assertEqual(owed.owed_on({"rows": [{"ingredient": "Limonene"}], "withheld": []}), [])

    def test_a_material_from_an_older_run_owes_nothing(self):
        """COMPOSITIONS.json written before any of this existed must not crash the workbook."""
        import owed

        self.assertEqual(owed.owed_on({"rows": [], "identifier": "R289"}), [])


if __name__ == "__main__":
    unittest.main()
